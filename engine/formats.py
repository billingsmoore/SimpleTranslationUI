"""File I/O: reading source documents (txt/docx/pdf) and writing translated
output (txt/docx/json). Segments are plain dicts: {"source": str, "target": str}.
"""

import json
from pathlib import Path

from .chunking import split_into_segments


def _extract_txt(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


def _extract_docx(path: str) -> str:
    import docx
    doc = docx.Document(path)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


_EXTRACTORS = {
    ".txt": _extract_txt,
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
}


def load_source_file(path: str, max_chars: int = 400) -> list[dict]:
    """Extract text from a .txt/.docx/.pdf file and split it into segments."""
    ext = Path(path).suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported file type: {ext}")
    text = extractor(path)
    if not text.strip():
        raise ValueError("No text could be extracted from this file.")
    chunks = split_into_segments(text, max_chars=max_chars)
    return [{"source": c, "target": ""} for c in chunks]


def parse_json(file_content: str | bytes) -> list[dict]:
    """Parse a previously-exported JSON file back into segments (for resume)."""
    if isinstance(file_content, bytes):
        file_content = file_content.decode("utf-8")
    data = json.loads(file_content)
    if not isinstance(data, list) or not data:
        raise ValueError("JSON must be a non-empty list of objects.")
    segments = []
    for i, item in enumerate(data):
        if not isinstance(item, dict) or "source" not in item:
            raise ValueError(f"Item {i} is missing 'source'.")
        segments.append({"source": item.get("source", ""), "target": item.get("target", "")})
    return segments


def export_to_txt(segments: list[dict], which: str = "target") -> str:
    """Join segment text with blank lines. which='target' or 'source'."""
    texts = [seg.get(which, "").strip() for seg in segments]
    texts = [t for t in texts if t]
    return "\n\n".join(texts) + "\n" if texts else ""


def export_to_docx(segments: list[dict], which: str = "target") -> bytes:
    import io
    import docx
    doc = docx.Document()
    for seg in segments:
        text = seg.get(which, "").strip()
        if text:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_to_json(segments: list[dict]) -> str:
    return json.dumps(
        [{"source": s.get("source", ""), "target": s.get("target", "")} for s in segments],
        ensure_ascii=False, indent=2,
    )
